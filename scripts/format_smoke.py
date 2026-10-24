"""Generate generic fixtures and verify text/table survival without hosted calls."""

import tempfile
from pathlib import Path

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from core.adapters.parsers.structured_adapter import StructuredParserAdapter
from core.chunking.strategy import split_documents


def build_fixtures(folder: Path) -> None:
    text = "The synthetic launch code is amber. The shipping limit is 42 units."
    (folder / "sample.txt").write_text(text)
    (folder / "sample.html").write_text(
        "<html><body><h1>Synthetic guide</h1><p>" + text + "</p><table>"
        "<tr><th>Item</th><th>Limit</th></tr><tr><td>shipping</td><td>42</td></tr>"
        "<tr><td>Synthetic item</td><td>73</td></tr></table></body></html>"
    )
    (folder / "sample.csv").write_text("Item,Limit\nshipping,42\nSynthetic item,73\n")
    doc = Document()
    doc.add_paragraph(text)
    table = doc.add_table(rows=3, cols=2)
    for row, values in enumerate((("Item", "Limit"), ("shipping", "42"), ("Synthetic item", "73"))):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    doc.save(folder / "sample.docx")
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    lines = [f"BT /F1 12 Tf 40 750 Td ({text}) Tj ET"]
    for y, values in ((700, ("Item", "Limit")), (680, ("shipping", "42")), (660, ("Synthetic item", "73"))):
        for x, value in zip((40, 220), values, strict=True):
            lines.append(f"BT /F1 12 Tf {x} {y} Td ({value}) Tj ET")
    for y in (715, 695, 675, 650):
        lines.append(f"35 {y} m 280 {y} l S")
    for x in (35, 210, 280):
        lines.append(f"{x} 650 m {x} 715 l S")
    stream = DecodedStreamObject()
    stream.set_data("\n".join(lines).encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    with (folder / "sample.pdf").open("wb") as output:
        writer.write(output)


def main() -> None:
    with tempfile.TemporaryDirectory() as directory:
        folder = Path(directory)
        build_fixtures(folder)
        for extension in ("pdf", "docx", "txt", "html", "csv"):
            documents = StructuredParserAdapter().parse(
                str(folder / ("sample." + extension)), source_id="synthetic", use_case="synthetic"
            )
            content = " ".join(doc.page_content for doc in split_documents(documents))
            assert "42" in content and "shipping" in content.lower(), extension
            if extension != "txt":
                assert "Synthetic item" in content and "73" in content, extension
            print(extension + ": text and applicable table values preserved")


if __name__ == "__main__":
    main()
