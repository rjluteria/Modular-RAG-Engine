"""Local text parsers preserving table cells without OCR or remote services."""

import csv
import io
from collections import Counter
from pathlib import Path

from bs4 import BeautifulSoup
from langchain_core.documents import Document

_SUPPORTED = {".pdf", ".docx", ".html", ".htm", ".txt", ".csv"}


def _text(value) -> str:
    return " ".join(str(value or "").split())


def _html_rows(table) -> tuple[list[list[str]], int]:
    grid: dict[tuple[int, int], str] = {}
    header_rows = 0
    rows = table.find_all("tr", recursive=True)
    for r, row in enumerate(rows):
        if row.find_parent("table") is not table:
            continue
        if row.find_parent("thead") or (r == header_rows and row.find("th")):
            header_rows = r + 1
        column = 0
        for cell in row.find_all(["td", "th"], recursive=False):
            while (r, column) in grid:
                column += 1
            height, width = int(cell.get("rowspan", 1)), int(cell.get("colspan", 1))
            if not 1 <= height <= 1000 or not 1 <= width <= 1000 or len(grid) + height * width > 100000:
                raise ValueError("Table dimensions exceed supported limits")
            for y in range(r, r + height):
                for x in range(column, column + width):
                    grid[y, x] = _text(cell.get_text(" ", strip=True))
            column += width
    width = max((x + 1 for _, x in grid), default=0)
    height = max((y + 1 for y, _ in grid), default=0)
    if width * height > 100000:
        raise ValueError("Table dimensions exceed supported limits")
    return [[grid.get((y, x), "") for x in range(width)] for y in range(height)], max(header_rows, 1)


def _table(rows, *, header_rows=1, context="", metadata=None, merged_title=False) -> Document | None:
    rows = [[_text(cell) for cell in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return None
    if sum(len(row) for row in rows) > 100000:
        raise ValueError("Table dimensions exceed supported limits")
    width = max(map(len, rows))
    if width * len(rows) > 100000:
        raise ValueError("Table dimensions exceed supported limits")
    rows = [row + [""] * (width - len(row)) for row in rows]
    # A merged title spanning the first row qualifies the table, not its columns.
    if merged_title and len(rows) > 1 and len({cell for cell in rows[0] if cell}) == 1 and width > 1:
        context = " ".join(filter(None, [context, next(cell for cell in rows[0] if cell)]))
        rows = rows[1:]
        header_rows = max(1, header_rows - 1)
    header_rows = min(header_rows, len(rows))
    headers = []
    for col in range(width):
        labels = list(dict.fromkeys(rows[r][col] for r in range(header_rows) if rows[r][col]))
        headers.append(" / ".join(labels) or f"Column {col + 1}")
    # Preserve repeated column names with a position rather than silently overwriting keys.
    counts = Counter(headers)
    headers = [f"{h} (column {i + 1})" if counts[h] > 1 else h for i, h in enumerate(headers)]
    data = [row for row in rows[header_rows:] if row not in rows[:header_rows]]
    return Document(
        page_content="\n".join(" | ".join(row) for row in rows),
        metadata={
            **(metadata or {}),
            "element_type": "table",
            "section_title": context or None,
            "table_headers": headers,
            "table_rows": data,
        },
    )


class StructuredParserAdapter:
    def supports(self, filename: str, content_type: str | None) -> bool:
        return Path(filename).suffix.lower() in _SUPPORTED

    def parse(self, file_path: str, *, source_id: str, use_case: str) -> list[Document]:
        path = Path(file_path)
        metadata = {
            "source_id": source_id,
            "use_case": use_case,
            "source_uri": file_path,
            "page_number": None,
            "section_title": None,
            "element_type": "text",
        }
        extension = path.suffix.lower()
        documents = []
        if extension == ".pdf":
            import pdfplumber

            with pdfplumber.open(path) as pdf:
                if len(pdf.pages) > 1000:
                    raise ValueError("PDF exceeds supported page limit")
                for number, page in enumerate(pdf.pages, 1):
                    page_meta = {**metadata, "page_number": number}
                    tables = page.find_tables()

                    def outside_tables(obj, tables=tables):
                        return not any(
                            t.bbox[0] <= obj.get("x0", -1)
                            and obj.get("x1", -1) <= t.bbox[2]
                            and t.bbox[1] <= obj.get("top", -1)
                            and obj.get("bottom", -1) <= t.bbox[3]
                            for t in tables
                        )

                    prose = page.filter(outside_tables).extract_text() or ""
                    if prose.strip():
                        documents.append(Document(page_content=prose, metadata=page_meta))
                    for pdf_table in tables:
                        pdf_rows = pdf_table.extract()
                        # Ruled merged header cells have None slots; retain the parent column label.
                        merged_title = bool(
                            pdf_rows
                            and len(pdf_rows[0]) > 1
                            and pdf_rows[0][0]
                            and all(c is None for c in pdf_rows[0][1:])
                        )
                        header_rows = (
                            2
                            if len(pdf_rows) > 2
                            and (
                                merged_title
                                or (
                                    any(c is None for c in pdf_rows[0])
                                    and any(c is None for c in pdf_rows[1])
                                )
                            )
                            else 1
                        )
                        if header_rows == 2:
                            previous = ""
                            for index, cell in enumerate(pdf_rows[0]):
                                previous = cell or previous
                                pdf_rows[0][index] = previous
                        caption = (
                            page.crop(
                                (0, max(0, pdf_table.bbox[1] - 60), page.width, pdf_table.bbox[1])
                            ).extract_text()
                            or ""
                        )
                        doc = _table(
                            pdf_rows,
                            header_rows=header_rows,
                            context=caption,
                            metadata=page_meta,
                            merged_title=merged_title,
                        )
                        if doc:
                            documents.append(doc)
        elif extension == ".docx":
            from docx import Document as WordDocument
            from docx.table import Table
            from docx.text.paragraph import Paragraph

            word = WordDocument(str(path))
            title = ""
            previous = ""
            for block in word.iter_inner_content():
                if isinstance(block, Paragraph):
                    text = block.text.strip()
                    if not text:
                        continue
                    if block.style and block.style.name and block.style.name.startswith("Heading"):
                        title = text
                    previous = text
                    documents.append(
                        Document(page_content=text, metadata={**metadata, "section_title": title or None})
                    )
                elif isinstance(block, Table):
                    rows = [[cell.text for cell in row.cells] for row in block.rows]
                    headers = 0
                    for row in block.rows:
                        if row._tr.xpath("./w:trPr/w:tblHeader"):
                            headers += 1
                        else:
                            break
                    doc = _table(
                        rows,
                        header_rows=max(headers, 1),
                        context=" ".join(dict.fromkeys(filter(None, [title, previous]))),
                        metadata=metadata,
                    )
                    if doc:
                        documents.append(doc)
        else:
            raw = path.read_bytes()
            text = raw.decode("utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig")
            if extension == ".csv":
                doc = _table(list(csv.reader(io.StringIO(text))), metadata=metadata)
                if doc:
                    documents.append(doc)
            elif extension in {".html", ".htm"}:
                soup = BeautifulSoup(text, "html.parser")
                for tag in soup(["script", "style", "noscript"]):
                    tag.decompose()
                for html_table in list(soup.find_all("table")):
                    if html_table.find_parent("table"):
                        continue
                    heading = html_table.find_previous(["h1", "h2", "h3", "h4", "p"])
                    html_caption = html_table.find("caption")
                    context = " ".join(t.get_text(" ", strip=True) for t in (heading, html_caption) if t)
                    html_rows, html_header_rows = _html_rows(html_table)
                    doc = _table(html_rows, header_rows=html_header_rows, context=context, metadata=metadata)
                    if doc:
                        documents.append(doc)
                    html_table.decompose()
                prose = soup.get_text("\n", strip=True)
                if prose:
                    documents.append(Document(page_content=prose, metadata=metadata))
            elif extension == ".txt":
                if text.strip():
                    documents.append(Document(page_content=text, metadata=metadata))
            else:
                raise ValueError("Unsupported document format")
        return documents
